"""Render a report written in Markdown as a PDF (reportlab) or a Word document (python-docx),
text and figures together, so a simulation report can be read and shared without a Markdown
viewer.

One parser serves both renderers. It covers the Markdown the report builder and an engineering
write-up use: headings, paragraphs (two trailing spaces or ``<br>`` break a line), ``-``/``*``
and numbered lists (nested by indentation), pipe tables, images on a line of their own (paths
relative to the Markdown file), fenced code, block quotes and rules; inline bold, italic, code
and links. Bold PASS / FAIL / INCOMPLETE are drawn in their status colours.

    from konnect_sim import document
    document.render("reports/buck_report.md", ["pdf", "docx"])
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

FORMATS = ("pdf", "docx")
STATUS_RGB = {"PASS": (0x1B, 0x7F, 0x4C), "FAIL": (0xC0, 0x30, 0x30), "INCOMPLETE": (0xA6, 0x6A, 0x00)}
INK, INK2, RULE, SHADE, ACCENT = (0x1F, 0x23, 0x28), (0x55, 0x5D, 0x66), (0xC9, 0xCE, 0xD4), (0xF1, 0xF3, 0xF5), \
    (0x2A, 0x5D, 0xA8)


# ── parsing ──────────────────────────────────────────────────────────────────

@dataclass
class Span:
    text: str
    bold: bool = False
    italic: bool = False
    code: bool = False
    link: str = ""
    brk: bool = False                 # a line break, text is empty


@dataclass
class Block:
    kind: str                         # heading, para, list, table, image, code, quote, rule
    spans: list[Span] = field(default_factory=list)
    level: int = 0                    # heading level; list nesting depth
    ordered: bool = False
    number: int = 0                   # numbered list item: its number
    rows: list[list[list[Span]]] = field(default_factory=list)   # table: header first
    text: str = ""                    # code; image alt
    path: str = ""                    # image file


_INLINE = re.compile(r"`([^`]+)`|\*\*(.+?)\*\*|__(.+?)__|(?<![\w*])\*(?![\s*])(.+?)(?<![\s*])\*(?![\w*])"
                     r"|(?<![\w_])_(?![\s_])(.+?)(?<![\s_])_(?![\w_])|\[([^\]]+)\]\(([^)\s]+)\)|<br\s*/?>")


def inline(text: str, bold=False, italic=False, link="") -> list[Span]:
    """Inline Markdown as styled spans."""
    out, pos = [], 0
    for m in _INLINE.finditer(text):
        if m.start() > pos:
            out.append(Span(text[pos:m.start()], bold, italic, link=link))
        code, b1, b2, i1, i2, ltext, url = m.groups()
        if code is not None:
            out.append(Span(code, bold, italic, code=True, link=link))
        elif b1 is not None or b2 is not None:
            out += inline(b1 if b1 is not None else b2, True, italic, link)
        elif i1 is not None or i2 is not None:
            out += inline(i1 if i1 is not None else i2, bold, True, link)
        elif ltext is not None:
            out += inline(ltext, bold, italic, url)
        else:
            out.append(Span("", brk=True))
        pos = m.end()
    if pos < len(text):
        out.append(Span(text[pos:], bold, italic, link=link))
    return [s for s in out if s.text or s.brk]


def _cells(line: str) -> list[str]:
    line = line.strip()
    line = line[1:] if line.startswith("|") else line
    line = line[:-1] if line.endswith("|") and not line.endswith("\\|") else line
    return [c.strip().replace("\\|", "|") for c in re.split(r"(?<!\\)\|", line)]


_IMAGE = re.compile(r"^!\[([^\]]*)\]\(([^)]+)\)\s*$")
_LIST = re.compile(r"^(\s*)([-*+]|\d+[.)])\s+(.*)$")
_TABLE_RULE = re.compile(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$")


def parse(md: str) -> list[Block]:
    lines = md.replace("\r\n", "\n").split("\n")
    blocks: list[Block] = []
    para: list[str] = []

    def flush():
        if para:
            text = ""
            for ln in para:
                hard = ln.endswith("  ")
                text += ln.strip() + ("<br>" if hard else " ")
            blocks.append(Block("para", inline(text.strip().removesuffix("<br>"))))
            para.clear()

    i = 0
    while i < len(lines):
        ln = lines[i]
        s = ln.strip()
        if not s:
            flush(); i += 1; continue
        if s.startswith("```"):
            flush()
            j = i + 1
            while j < len(lines) and not lines[j].strip().startswith("```"):
                j += 1
            blocks.append(Block("code", text="\n".join(lines[i + 1:j])))
            i = j + 1; continue
        if m := re.match(r"^(#{1,6})\s+(.*?)\s*#*$", s):
            flush()
            blocks.append(Block("heading", inline(m.group(2)), level=len(m.group(1))))
            i += 1; continue
        if re.match(r"^(-{3,}|\*{3,}|_{3,})$", s):
            flush(); blocks.append(Block("rule")); i += 1; continue
        if m := _IMAGE.match(s):
            flush()
            blocks.append(Block("image", text=m.group(1), path=m.group(2).strip().strip("<>")))
            i += 1; continue
        if s.startswith("|") and i + 1 < len(lines) and _TABLE_RULE.match(lines[i + 1]):
            flush()
            rows = [_cells(s)]
            j = i + 2
            while j < len(lines) and lines[j].strip().startswith("|"):
                rows.append(_cells(lines[j]))
                j += 1
            width = len(rows[0])
            rows = [(r + [""] * width)[:width] for r in rows]
            blocks.append(Block("table", rows=[[inline(c) for c in r] for r in rows]))
            i = j; continue
        if s.startswith(">"):
            flush()
            q = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                q.append(lines[i].strip()[1:].strip())
                i += 1
            blocks.append(Block("quote", inline(" ".join(q))))
            continue
        if m := _LIST.match(ln):
            flush()
            indent, marker, text = m.groups()
            j = i + 1
            while j < len(lines) and lines[j].strip() and not _LIST.match(lines[j]) \
                    and lines[j].startswith(" ") and not lines[j].strip().startswith("|"):
                text += " " + lines[j].strip()
                j += 1
            ordered = marker[0].isdigit()
            blocks.append(Block("list", inline(text), level=len(indent.expandtabs(4)) // 2, ordered=ordered,
                                number=int(marker[:-1]) if ordered else 0))
            i = j; continue
        para.append(ln)
        i += 1
    flush()
    return blocks


def plain(spans: list[Span]) -> str:
    return "".join("\n" if s.brk else s.text for s in spans)


# ── PDF ──────────────────────────────────────────────────────────────────────

def _fonts() -> dict[str, str]:
    """DejaVu (shipped with matplotlib) covers µ, Ω, ≤, ±, °, → and Greek."""
    from matplotlib import get_data_path
    from reportlab.lib.fonts import addMapping
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    d = Path(get_data_path()) / "fonts" / "ttf"
    faces = {"Body": "DejaVuSans.ttf", "Body-Bold": "DejaVuSans-Bold.ttf", "Body-Italic": "DejaVuSans-Oblique.ttf",
             "Body-BoldItalic": "DejaVuSans-BoldOblique.ttf", "Mono": "DejaVuSansMono.ttf"}
    for name, file in faces.items():
        if name not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont(name, str(d / file)))
    addMapping("Body", 0, 0, "Body"); addMapping("Body", 1, 0, "Body-Bold")
    addMapping("Body", 0, 1, "Body-Italic"); addMapping("Body", 1, 1, "Body-BoldItalic")
    return faces


def _hex(rgb) -> str:
    return "#%02x%02x%02x" % rgb


def _markup(spans: list[Span]) -> str:
    from xml.sax.saxutils import escape
    out = []
    for s in spans:
        if s.brk:
            out.append("<br/>"); continue
        t = escape(s.text)
        if s.code:
            t = f'<font name="Mono" backColor="{_hex(SHADE)}">{t}</font>'
        if s.bold and s.text.strip() in STATUS_RGB:
            t = f'<font color="{_hex(STATUS_RGB[s.text.strip()])}">{t}</font>'
        if s.italic:
            t = f"<i>{t}</i>"
        if s.bold:
            t = f"<b>{t}</b>"
        if s.link:
            t = f'<link href="{escape(s.link)}" color="{_hex(ACCENT)}">{t}</link>'
        out.append(t)
    return "".join(out)


def _col_widths(rows: list[list[str]], total: float, measure=lambda t: 5.2 * len(t)) -> list[float]:
    """Share the frame width by content length, so short columns stay narrow and long text wraps;
    no column is narrower than its longest word (``measure``: text -> points)."""
    n, pad = len(rows[0]), 8.0
    floor = [max([measure(w) for r in rows for w in r[c].split()] + [measure("mm")]) + pad for c in range(n)]
    need = [min(max(sum(len(r[c]) for r in rows) / len(rows), 6.0), 70.0) for c in range(n)]
    if sum(floor) >= total:                         # words alone overflow: shrink in proportion
        return [total * f / sum(floor) for f in floor]
    share = total - sum(floor)                      # the rest goes by need, on top of each floor
    return [floor[c] + share * need[c] / sum(need) for c in range(n)]


def to_pdf(md_path: str | Path, pdf_path: str | Path | None = None) -> Path:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.pdfbase.pdfmetrics import stringWidth
    from reportlab.platypus import (CondPageBreak, HRFlowable, Image, KeepTogether, ListFlowable, ListItem,
                                    Paragraph, Preformatted, SimpleDocTemplate, Spacer, Table, TableStyle)
    from PIL import Image as PILImage

    md_path = Path(md_path)
    pdf_path = Path(pdf_path) if pdf_path else md_path.with_suffix(".pdf")
    _fonts()
    blocks = parse(md_path.read_text(encoding="utf-8"))
    title = next((plain(b.spans) for b in blocks if b.kind == "heading" and b.level == 1), md_path.stem)
    ink, ink2 = colors.Color(*[x / 255 for x in INK]), colors.Color(*[x / 255 for x in INK2])
    base = ParagraphStyle("base", fontName="Body", fontSize=9.5, leading=13.2, textColor=ink, spaceAfter=5)
    st = {
        1: ParagraphStyle("h1", base, fontName="Body-Bold", fontSize=17, leading=21, spaceAfter=8),
        2: ParagraphStyle("h2", base, fontName="Body-Bold", fontSize=13, leading=17, spaceBefore=12, spaceAfter=5, keepWithNext=1),
        3: ParagraphStyle("h3", base, fontName="Body-Bold", fontSize=10.5, leading=14, spaceBefore=8, spaceAfter=3, keepWithNext=1),
        4: ParagraphStyle("h4", base, fontName="Body-Bold", fontSize=9.5, leading=13, spaceBefore=6, spaceAfter=2, keepWithNext=1),
    }
    cell = ParagraphStyle("cell", base, fontSize=7.6, leading=9.6, spaceAfter=0)
    head = ParagraphStyle("head", cell, fontName="Body-Bold")
    caption = ParagraphStyle("cap", base, fontSize=8, leading=10.5, textColor=ink2, alignment=TA_CENTER,
                             spaceBefore=2, spaceAfter=10)
    quote = ParagraphStyle("quote", base, leftIndent=12, textColor=ink2)
    mono = ParagraphStyle("mono", base, fontName="Mono", fontSize=7.6, leading=9.6, backColor=colors.Color(
        *[x / 255 for x in SHADE]), borderPadding=5, leftIndent=5, rightIndent=5, spaceBefore=3, spaceAfter=8)

    doc = SimpleDocTemplate(str(pdf_path), pagesize=letter, leftMargin=0.8 * inch, rightMargin=0.8 * inch,
                            topMargin=0.75 * inch, bottomMargin=0.75 * inch, title=title, author="konnect_sim")
    fw, fh = doc.width, doc.height
    story, pending_list, figure_no, heading = [], [], 0, None

    def flush_list():
        if not pending_list:
            return
        def build(items, depth):
            out, k = [], 0
            while k < len(items):
                b = items[k]
                sub = []
                k += 1
                while k < len(items) and items[k].level > depth:
                    sub.append(items[k]); k += 1
                flow = [Paragraph(_markup(b.spans), ParagraphStyle("li", base, spaceAfter=2))]
                if sub:
                    flow.append(build(sub, depth + 1))
                out.append(ListItem(flow))
            kw = dict(leftIndent=14, bulletFontName="Body", bulletColor=ink2)
            if items[0].ordered:
                return ListFlowable(out, bulletType="1", start=items[0].number or 1, bulletFontSize=9, **kw)
            return ListFlowable(out, bulletType="bullet", start="•", bulletFontSize=8, **kw)
        depth0 = min(b.level for b in pending_list)
        for b in pending_list:
            b.level -= depth0
        story.append(build(list(pending_list), 0))
        story.append(Spacer(1, 4))
        pending_list.clear()

    for b in blocks:
        if b.kind != "list":
            flush_list()
        if b.kind == "heading":
            if b.level == 2:
                story.append(CondPageBreak(1.4 * inch))
            heading = Paragraph(_markup(b.spans), st[min(b.level, 4)])
            story.append(heading)
            if b.level == 1:
                story.append(HRFlowable(width="100%", thickness=1.2, color=colors.Color(*[x / 255 for x in ACCENT]),
                                        spaceAfter=8))
        elif b.kind == "para":
            story.append(Paragraph(_markup(b.spans), base))
        elif b.kind == "list":
            pending_list.append(b)
        elif b.kind == "quote":
            story.append(Paragraph(_markup(b.spans), quote))
        elif b.kind == "rule":
            story.append(HRFlowable(width="100%", thickness=0.6, color=colors.Color(*[x / 255 for x in RULE]),
                                    spaceBefore=4, spaceAfter=8))
        elif b.kind == "code":
            story.append(Preformatted(b.text, mono))
        elif b.kind == "table":
            data = [[Paragraph(_markup(c), head if r == 0 else cell) for c in row] for r, row in enumerate(b.rows)]
            widths = _col_widths([[plain(c) for c in row] for row in b.rows], fw,
                                 lambda t: stringWidth(t, "Body-Bold", cell.fontSize))
            t = Table(data, colWidths=widths, repeatRows=1, splitInRow=1)
            t.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.Color(*[x / 255 for x in SHADE])),
                ("LINEBELOW", (0, 0), (-1, 0), 0.8, ink2),
                ("LINEBELOW", (0, 1), (-1, -1), 0.3, colors.Color(*[x / 255 for x in RULE])),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 3), ("RIGHTPADDING", (0, 0), (-1, -1), 3),
                ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2)]))
            story += [t, Spacer(1, 8)]
        elif b.kind == "image":
            p = Path(b.path)
            p = p if p.is_absolute() else (md_path.parent / p)
            if not p.is_file():
                story.append(Paragraph(f"<i>[missing figure: {b.path}]</i>", quote))
                continue
            with PILImage.open(p) as im:
                w, h = im.size
            scale = min(fw / w, 0.62 * fh / h, 72 / 150)     # never enlarged past 150 dpi
            figure_no += 1
            lead = [story.pop()] if story and story[-1] is heading else []   # a figure's heading stays with it
            story.append(KeepTogether(lead + [Image(str(p), width=w * scale, height=h * scale),
                                              Paragraph(f"Figure {figure_no}. {_markup(inline(b.text))}", caption)]))
    flush_list()

    def frame(canvas, d):
        canvas.saveState()
        canvas.setFont("Body", 7.5)
        canvas.setFillColor(ink2)
        canvas.drawString(d.leftMargin, 0.45 * inch, title[:110])
        canvas.drawRightString(d.leftMargin + d.width, 0.45 * inch, f"page {d.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=frame, onLaterPages=frame)
    return pdf_path


# ── Word ─────────────────────────────────────────────────────────────────────

def _docx_runs(par, spans: list[Span], size=None):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    def style(run, s):
        run.bold, run.italic = s.bold or None, s.italic or None
        if s.code:
            run.font.name = "Consolas"
        if size:
            run.font.size = Pt(size)
        if s.bold and s.text.strip() in STATUS_RGB:
            run.font.color.rgb = RGBColor(*STATUS_RGB[s.text.strip()])

    for s in spans:
        if s.brk:
            par.add_run().add_break(); continue
        if s.link:
            rid = par.part.relate_to(s.link, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/"
                                     "hyperlink", is_external=True)
            h = OxmlElement("w:hyperlink"); h.set(qn("r:id"), rid)
            run = par.add_run(s.text)
            style(run, s)
            run.font.color.rgb = RGBColor(*ACCENT); run.font.underline = True
            h.append(run._r)
            par._p.append(h)
            continue
        style(par.add_run(s.text), s)


def to_docx(md_path: str | Path, docx_path: str | Path | None = None) -> Path:
    import docx
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
    from PIL import Image as PILImage

    md_path = Path(md_path)
    docx_path = Path(docx_path) if docx_path else md_path.with_suffix(".docx")
    blocks = parse(md_path.read_text(encoding="utf-8"))
    d = docx.Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = Inches(0.8)
    sec.top_margin = sec.bottom_margin = Inches(0.75)
    width = sec.page_width - sec.left_margin - sec.right_margin
    normal = d.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(10)
    normal.paragraph_format.space_after = Pt(4)

    foot = sec.footer.paragraphs[0]
    foot.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = foot.add_run("page "); r.font.size = Pt(8)
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        run = foot.add_run(); run.font.size = Pt(8)
        if tag:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = text
        run._r.append(el)

    figure_no = 0
    for b in blocks:
        if b.kind == "heading":
            p = d.add_heading(level=0 if b.level == 1 else min(b.level - 1, 4))
            _docx_runs(p, b.spans)
        elif b.kind == "para":
            _docx_runs(d.add_paragraph(), b.spans)
        elif b.kind == "list":
            name = ("List Number" if b.ordered else "List Bullet") + (f" {min(b.level + 1, 3)}" if b.level else "")
            _docx_runs(d.add_paragraph(style=name), b.spans)
        elif b.kind == "quote":
            p = d.add_paragraph(style="Quote")
            _docx_runs(p, b.spans)
        elif b.kind == "rule":
            p = d.add_paragraph()
            bdr = OxmlElement("w:pBdr"); bot = OxmlElement("w:bottom")
            for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "C9CED4")):
                bot.set(qn(k), v)
            bdr.append(bot); p._p.get_or_add_pPr().append(bdr)
        elif b.kind == "code":
            p = d.add_paragraph()
            run = p.add_run(b.text); run.font.name = "Consolas"; run.font.size = Pt(8)
        elif b.kind == "table":
            t = d.add_table(rows=len(b.rows), cols=len(b.rows[0]))
            t.style = "Table Grid"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            widths = _col_widths([[plain(c) for c in row] for row in b.rows], float(width))
            for ri, row in enumerate(b.rows):
                for ci, spans in enumerate(row):
                    c = t.cell(ri, ci)
                    c.width = int(widths[ci])
                    par = c.paragraphs[0]
                    par.paragraph_format.space_after = Pt(0)
                    _docx_runs(par, [Span(s.text, s.bold or ri == 0, s.italic, s.code, s.link, s.brk) for s in spans],
                               size=8)
                    if ri == 0:
                        shd = OxmlElement("w:shd")
                        for k, v in (("w:val", "clear"), ("w:color", "auto"), ("w:fill", "F1F3F5")):
                            shd.set(qn(k), v)
                        c._tc.get_or_add_tcPr().append(shd)
            hdr = t.rows[0]._tr.get_or_add_trPr()
            rep = OxmlElement("w:tblHeader"); rep.set(qn("w:val"), "true"); hdr.append(rep)
            d.add_paragraph()
        elif b.kind == "image":
            p = Path(b.path)
            p = p if p.is_absolute() else (md_path.parent / p)
            if not p.is_file():
                d.add_paragraph(f"[missing figure: {b.path}]").runs[0].italic = True
                continue
            with PILImage.open(p) as im:
                w, h = im.size
            max_h = Inches(5.8)
            pic_w = min(int(width), int(max_h * w / h))
            d.add_picture(str(p), width=pic_w)
            d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            figure_no += 1
            cap = d.add_paragraph(style="Caption")
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _docx_runs(cap, [Span(f"Figure {figure_no}. ")] + inline(b.text))
            for run in cap.runs:
                run.font.color.rgb = RGBColor(*INK2)
    d.core_properties.title = next((plain(b.spans) for b in blocks if b.kind == "heading" and b.level == 1),
                                   md_path.stem)
    d.save(str(docx_path))
    return docx_path


def render(md_path: str | Path, formats=("pdf",)) -> dict[str, str]:
    """Write ``<md stem>.pdf`` / ``.docx`` next to the Markdown file; returns {format: path}."""
    bad = [f for f in formats if f not in FORMATS]
    if bad:
        raise ValueError(f"unknown document format(s) {bad}; use {list(FORMATS)}")
    out = {}
    for f in formats:
        out[f] = str((to_pdf if f == "pdf" else to_docx)(md_path))
    return out
